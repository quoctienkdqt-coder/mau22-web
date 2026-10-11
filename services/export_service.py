import io
import re
import copy
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from vessel_utils import deduplicate_vessels

class ExportService:
    @staticmethod
    def _clean_text(text):
        """Làm sạch ký tự lạ để tránh lỗi XML khi tạo Word"""
        if text is None:
            return ""
        s = str(text).strip()
        if s.lower() in ("nan", "none", "null"):
            return ""
        return re.sub(r'[^\x09\x0A\x0D\x20-\x7E\x85\xA0-\uD7FF\uE000-\uFFFD\U00010000-\U0010ffff]', '', s).strip()

    @staticmethod
    def export_mau22(data, save_path=None):
        """
        Tạo file Word Mẫu 22.KT từ dict data.
        Hỗ trợ:
        - save_path=None: trả về bytes trong bộ nhớ
        - save_path là str hoặc Path: lưu ra file đĩa và trả về đường dẫn
        - save_path là stream (io.BytesIO): ghi trực tiếp vào luồng và trả về stream
        """
        data_copy = copy.deepcopy(data) if isinstance(data, dict) else {}
        if "tau_danh_bat" in data_copy and isinstance(data_copy["tau_danh_bat"], list):
            data_copy["tau_danh_bat"] = deduplicate_vessels(data_copy["tau_danh_bat"])

        doc = Document()
        sec = doc.sections[0]
        sec.orientation = WD_ORIENT.LANDSCAPE
        if sec.page_width < sec.page_height:
            sec.page_width, sec.page_height = sec.page_height, sec.page_width
        
        # Cài đặt lề trang chuẩn 1.5 cm trên cả 4 cạnh
        sec.left_margin = Cm(1.5)
        sec.right_margin = Cm(1.5)
        sec.top_margin = Cm(1.5)
        sec.bottom_margin = Cm(1.5)
        
        doc.styles['Normal'].font.name = 'Times New Roman'
        doc.styles['Normal'].font.size = Pt(11)

        # Header hành chính
        doc.add_paragraph("Mẫu 22.KT").alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header_table = doc.add_table(rows=1, cols=2)
        header_table.autofit = False
        header_table.columns[0].width = Cm(12)
        header_table.columns[1].width = Cm(14)

        # Cột 1
        p1 = header_table.cell(0, 0).paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.add_run(f"{ExportService._clean_text(data_copy.get('ten_co_so', '')).upper()}\n").bold = True
        p1.add_run("───────\nSố/No: ............")

        # Cột 2
        p2 = header_table.cell(0, 1).paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập - Tự do - Hạnh phúc\n").bold = True
        p2.add_run("SOCIALIST REPUBLIC OF VIETNAM\nIndependence - Freedom - Happiness\n").bold = True
        p2.add_run("────────")
        
        doc.add_paragraph()

        # Thông tin cơ sở
        p_kinh_gui = doc.add_paragraph(f"Kính gửi/To: {ExportService._clean_text(data_copy.get('kinh_gui', ''))}")
        p_kinh_gui.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if p_kinh_gui.runs:
            p_kinh_gui.runs[0].bold = True
        
        doc.add_paragraph(f"Tên cơ sở nhập khẩu/Importer: {ExportService._clean_text(data_copy.get('ten_co_so', ''))}")
        doc.add_paragraph(f"Địa chỉ/Address: {ExportService._clean_text(data_copy.get('dia_chi', ''))}")
        doc.add_paragraph(f"Người đại diện/Representative: {ExportService._clean_text(data_copy.get('nguoi_dai_dien', ''))}")
        doc.add_paragraph(f"Số điện thoại/Tel: {ExportService._clean_text(data_copy.get('sdt', ''))}")
        doc.add_paragraph("\nChúng tôi khai báo thủy sản, sản phẩm thủy sản nhập khẩu từ các tàu công ten nơ vào Việt Nam như sau:")

        # Helper vẽ bảng nghiệp vụ chuẩn
        def add_table(title, headers, dict_key, dict_keys_ordered):
            p_title = doc.add_paragraph(f"\n{title}")
            if p_title.runs:
                p_title.runs[0].bold = True
            if dict_key == "san_pham":
                doc.add_paragraph(f"- Tổng sản lượng cá sẽ được bốc dỡ/Catch to be offloaded: {ExportService._clean_text(data_copy.get('tong_san_luong', ''))} kg")
            
            t = doc.add_table(rows=1, cols=len(headers))
            t.style = 'Table Grid'

            # Header row định dạng in đậm và căn giữa
            hdr_cells = t.rows[0].cells
            for i, h in enumerate(headers):
                hdr_cells[i].text = h
                p = hdr_cells[i].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if p.runs:
                    p.runs[0].bold = True

            rows_data = data_copy.get(dict_key, [])
            if not isinstance(rows_data, list):
                rows_data = []

            for idx, row in enumerate(rows_data):
                if not isinstance(row, dict):
                    continue
                cells = t.add_row().cells
                cells[0].text = str(idx + 1)
                cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                for i, k in enumerate(dict_keys_ordered):
                    cells[i + 1].text = ExportService._clean_text(row.get(k, ''))

        # 6 Bảng chính quy
        add_table("1. Thủy sản, sản phẩm thủy sản nhập khẩu/Imported catch products", ["TT", "Loài", "Mã HS", "Tên KH", "Kích cỡ", "Khối lượng", "Xuất xứ"], "san_pham", ["loai", "hs_code", "ten_khoa_hoc", "kich_co", "khoi_luong", "xuat_xu"])
        doc.add_paragraph()
        doc.add_paragraph(f" - Thời gian bốc dỡ/Offloading Time: {ExportService._clean_text(data_copy.get('thoi_gian_boc_do', ''))}")
        doc.add_paragraph(f" - Địa điểm bốc dỡ/Offloading Venue: {ExportService._clean_text(data_copy.get('dia_diem_boc_do', ''))}")
        add_table("2. Thông tin về tàu đánh bắt/Fishing Vessel(s) information", ["TT", "Tên tàu", "IMO", "Hô hiệu", "RFMO", "Giấy phép", "Cơ quan", "Hết hạn", "Khu vực", "Đối tượng", "Ngư cụ"], "tau_danh_bat", ["ten", "imo", "ho_hieu", "rfmo", "giay_phep", "co_quan", "ngay_het_han", "khu_vuc", "doi_tuong", "ngu_cu"])
        add_table("3. Thông tin chuyển tải liên quan (Nếu có)/Transshipment information (If applicable)", ["TT", "Tàu khai thác", "Thời gian", "Tàu nhận", "Cờ", "IMO", "Hô hiệu", "Loài", "Khối lượng"], "thong_tin_chuyen_tai", ["tau_khai_thac", "thoi_gian", "tau_nhan", "quoc_gia", "imo", "ho_hieu", "loai", "khoi_luong"])
        add_table("4. Các giấy phép chuyển tải có liên quan/Relevant transshipment authorization(s):", ["TT", "Tàu chuyển tải", "IMO", "Giấy phép", "Cơ quan", "Hết hạn"], "giay_phep_chuyen_tai", ["tau_chuyen_tai", "imo", "giay_phep", "co_quan", "ngay_het_han"])
        add_table("5. Thông tin về sản lượng khai thác cập bến/Landing information", ["TT", "Tên tàu", "IMO", "Cảng", "Ngày cập bến", "Khối lượng"], "san_luong_cap_ben", ["tau", "imo", "cang", "ngay", "khoi_luong"])
        add_table("6. Thông tin về công ten nơ", ["TT", "Container", "Seal", "Vận đơn", "Tàu chở", "IMO", "Nước xuất", "Cảng xuất", "Kho"], "cong_ten_no", ["so_container", "so_seal", "so_bl", "tau_cho", "imo", "nuoc_xuat", "cang_xuat", "dia_chi_kho"])
        
        doc.add_paragraph()
        p_cam_ket = doc.add_paragraph("Chúng tôi cam kết: Lô hàng có thông tin nêu trên không vi phạm IUU và hoàn toàn chịu trách nhiệm trước pháp luật trong trường hợp lô hàng vi phạm IUU; các số liệu khai báo trên là đúng sự thật, các hồ sơ tài liệu gửi kèm là hợp pháp, nếu sai chúng tôi hoàn toàn chịu trách nhiệm trước pháp luật/")
        p_cam_ket.add_run("We undertake that the consignment, including the above information, does not contain IUU products and shall assume full responsibility before the law for any IUU violations found; The above declared information is accurate, the attached documents and materials are lawful and valid, if it is shown to be inaccurate, we will be fully responsible before the law.").italic = True
        p_cam_ket.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        
        doc.add_paragraph()
        sig_table = doc.add_table(rows=1, cols=2)
        sig_table.autofit = False
        sig_table.columns[0].width = Cm(15) 
        sig_table.columns[1].width = Cm(10)
        
        p_sig = sig_table.cell(0, 1).paragraphs[0]
        p_sig.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sig.add_run("CHỦ CƠ SỞ NHẬP KHẨU/IMPORTER\n").bold = True
        p_sig.add_run("(Ký tên và đóng dấu/Signed and sealed)").italic = True

        if save_path is None:
            mem_stream = io.BytesIO()
            doc.save(mem_stream)
            return mem_stream.getvalue()
        elif hasattr(save_path, "write"):
            doc.save(save_path)
            return save_path
        else:
            doc.save(save_path)
            return save_path

    @staticmethod
    def export_mau22_bytes(data) -> bytes:
        """Xuất Mẫu 22.KT trực tiếp ra chuỗi byte trong RAM"""
        return ExportService.export_mau22(data, save_path=None)